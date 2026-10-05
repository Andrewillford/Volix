"""
Mapa de Ocorrências — a lente de QUALIDADE / não-conformidade sobre as entregas.

Onde (região/UF) e com quem (transportadora) os problemas se concentram, e quanto
custam. Nada é capturado de novo: cada ocorrência já foi importada pela Torre de
Controle (EDI OCOREN, tabela `ocorrencias` em tracking.db, ligada ao CTe por
`chave_cte`) e o valor da carga vem do histórico de auditoria (app/history.py,
`valor_mercadoria` do vCarga do CTe, mesma `chave_cte`).

Universo (denominador) = entregas auditadas do período (uma por CTe). Para cada
entrega olhamos as ocorrências dela e classificamos pela categoria já mapeada em
tables/ocorrencias_codigos.json (SINISTRO / DEVOLUCAO / TENTATIVA_FRUSTRADA /
EM_TRANSITO / ENTREGUE / OUTROS). As três primeiras são a NÃO-CONFORMIDADE — o
foco do painel; as informativas ficam disponíveis num toggle.

Métricas (quanto MENOR, melhor — como o %NF, ao contrário do OTD):
  • Taxa de ocorrência = entregas afetadas ÷ entregas totais × 100
    (afetada = teve ≥1 ocorrência problemática; contada UMA vez, não por evento).
  • Taxa de sinistro / devolução = idem, por categoria.
  • R$ sob risco = Σ valor da carga das entregas afetadas (por categoria conta o
    valor da entrega uma vez em cada categoria que ela tiver; no total, uma vez só).

Metas por categoria (alvo/atenção), editáveis no painel e guardadas em
tables/ocorrencias_metas.json — coloridas verde/amarelo/vermelho.
"""
from __future__ import annotations

import json
import os
from pathlib import Path
from typing import Dict, List, Optional

from app import history as audit_history
from app import tracking

METAS_FILE = os.getenv("OCORRENCIAS_METAS_FILE", str(Path(__file__).parent.parent / "tables" / "ocorrencias_metas.json"))

# Categorias problemáticas (a não-conformidade) e o rótulo de cada uma.
PROBLEMA = ("SINISTRO", "DEVOLUCAO", "TENTATIVA_FRUSTRADA")
CATEGORIAS = ("SINISTRO", "DEVOLUCAO", "TENTATIVA_FRUSTRADA", "EM_TRANSITO", "ENTREGUE", "OUTROS")
CAT_LABEL = {
    "SINISTRO": "Sinistro (extravio/avaria/roubo)",
    "DEVOLUCAO": "Devolução",
    "TENTATIVA_FRUSTRADA": "Tentativa frustrada",
    "EM_TRANSITO": "Em trânsito (informativa)",
    "ENTREGUE": "Entregue (informativa)",
    "OUTROS": "Outros",
}

# Metas padrão por categoria (% de entregas afetadas; MENOR é melhor). Placeholders
# conservadores — o usuário calibra no painel quando houver base real.
_META_SEED = {
    "taxa_ocorrencia":     {"meta": 5.0,  "atencao": 10.0},  # qualquer problema
    "sinistro":            {"meta": 0.5,  "atencao": 1.5},
    "devolucao":           {"meta": 2.0,  "atencao": 5.0},
    "tentativa_frustrada": {"meta": 3.0,  "atencao": 7.0},
}
_META_KEYS = tuple(_META_SEED.keys())
# De-para categoria problemática → chave de meta.
_CAT_META = {"SINISTRO": "sinistro", "DEVOLUCAO": "devolucao", "TENTATIVA_FRUSTRADA": "tentativa_frustrada"}

REGIOES = {
    "Norte":        ["AC", "AP", "AM", "PA", "RO", "RR", "TO"],
    "Nordeste":     ["AL", "BA", "CE", "MA", "PB", "PE", "PI", "RN", "SE"],
    "Centro-Oeste": ["DF", "GO", "MT", "MS"],
    "Sudeste":      ["ES", "MG", "RJ", "SP"],
    "Sul":          ["PR", "RS", "SC"],
}
UFS = {uf for ufs in REGIOES.values() for uf in ufs}

_metas_cache: Optional[Dict] = None


def _seed_metas() -> Dict:
    return {k: dict(v) for k, v in _META_SEED.items()}


def load_metas(force: bool = False) -> Dict:
    """Metas por categoria. Cria o arquivo com sementes no 1º uso; falha de
    leitura/escrita cai nas sementes em memória, sem quebrar o painel."""
    global _metas_cache
    if _metas_cache is not None and not force:
        return _metas_cache
    p = Path(METAS_FILE)
    data = None
    if p.exists():
        try:
            data = json.loads(p.read_text(encoding="utf-8"))
        except Exception:
            data = None
    if data is None:
        data = _seed_metas()
        try:
            p.parent.mkdir(parents=True, exist_ok=True)
            p.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        except Exception:
            pass
    # Garante todas as chaves esperadas (arquivo antigo/incompleto não quebra).
    for k, v in _META_SEED.items():
        data.setdefault(k, dict(v))
    _metas_cache = data
    return data


def save_metas(metas: Dict) -> Dict:
    """Grava as metas por categoria e atualiza o cache. Atenção nunca abaixo do alvo."""
    def _pair(d: Dict, seed: Dict) -> Dict:
        meta = round(float(d.get("meta", seed["meta"])), 2)
        atencao = round(float(d.get("atencao", seed["atencao"])), 2)
        if atencao < meta:
            atencao = meta
        return {"meta": meta, "atencao": atencao}

    data = {k: _pair((metas or {}).get(k, {}), _META_SEED[k]) for k in _META_KEYS}
    Path(METAS_FILE).parent.mkdir(parents=True, exist_ok=True)
    Path(METAS_FILE).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    global _metas_cache
    _metas_cache = data
    return data


def status_taxa(pct: Optional[float], chave_meta: str) -> Optional[str]:
    """'verde'/'amarelo'/'vermelho' pela faixa da taxa (MENOR = melhor) contra a
    meta da categoria. None quando não há base p/ calcular."""
    if pct is None:
        return None
    m = load_metas().get(chave_meta) or _META_SEED[chave_meta]
    meta, atencao = float(m["meta"]), float(m["atencao"])
    if pct <= meta:
        return "verde"
    if pct <= atencao:
        return "amarelo"
    return "vermelho"


def _uf_de(destino: Optional[str]) -> Optional[str]:
    """UF a partir do campo destino ('Município/UF')."""
    d = destino or ""
    if "/" in d:
        uf = d.rsplit("/", 1)[-1].strip().upper()
        if uf in UFS:
            return uf
    return None


# ─── Montagem do universo enriquecido ────────────────────────────────────────

def _universo(mes: Optional[str]) -> List[Dict]:
    """Uma linha por entrega auditada do período, já com transportadora, UF, valor
    da carga e as ocorrências dela agrupadas por categoria. `mes` = 'YYYY-MM'
    (mês de referência do CTe, mesma base dos demais painéis)."""
    ocs_por_chave = tracking.ocorrencias_por_chave()
    universo = []
    for r in audit_history.query_history(mes=mes, limit=100000):
        if r.get("status") == "ERRO_LEITURA":
            continue  # CTe ilegível não é uma entrega válida
        chave = r.get("chave_cte")
        ocs = ocs_por_chave.get(chave, []) if chave else []
        cats = {o.get("categoria") or "OUTROS" for o in ocs}
        universo.append({
            "chave": chave,
            "numero_cte": r.get("numero_cte"),
            "transportadora": (r.get("transportadora_tabela") or r.get("transportadora_nome_cte") or "Desconhecida").strip() or "Desconhecida",
            "uf": _uf_de(r.get("destino")),
            "destino": r.get("destino"),
            "valor": r.get("valor_mercadoria") or 0.0,
            "mes": r.get("mes_referencia") or (r.get("data_emissao") or "")[:7],
            "ocorrencias": ocs,
            "categorias": cats,
        })
    return universo


def _agrega(rows: List[Dict]) -> Dict:
    """Consolida uma lista de entregas (do _universo) nas métricas de ocorrência.
    Taxa = entregas afetadas ÷ entregas totais; R$ sob risco = valor das afetadas.
    por_categoria traz, para CADA categoria, entregas/ocorrências/valor."""
    total = len(rows)
    por_cat = {c: {"entregas": 0, "ocorrencias": 0, "valor": 0.0} for c in CATEGORIAS}
    afetadas = 0           # entregas com ≥1 ocorrência problemática
    valor_risco = 0.0      # valor das entregas afetadas (uma vez cada)
    ocorrencias_total = 0  # eventos problemáticos

    for r in rows:
        cats = r["categorias"]
        valor = r["valor"] or 0.0
        # Contagem de eventos por categoria.
        for o in r["ocorrencias"]:
            c = o.get("categoria") or "OUTROS"
            if c in por_cat:
                por_cat[c]["ocorrencias"] += 1
                if c in PROBLEMA:
                    ocorrencias_total += 1
        # Entregas afetadas e valor por categoria (uma vez por entrega).
        for c in cats:
            if c in por_cat:
                por_cat[c]["entregas"] += 1
                por_cat[c]["valor"] += valor
        if cats & set(PROBLEMA):
            afetadas += 1
            valor_risco += valor

    def _taxa(afet: int) -> Optional[float]:
        return round(afet / total * 100, 2) if total else None

    taxa = _taxa(afetadas)
    taxa_sin = _taxa(por_cat["SINISTRO"]["entregas"])
    taxa_dev = _taxa(por_cat["DEVOLUCAO"]["entregas"])
    taxa_tf = _taxa(por_cat["TENTATIVA_FRUSTRADA"]["entregas"])

    for c in CATEGORIAS:
        por_cat[c]["valor"] = round(por_cat[c]["valor"], 2)
        por_cat[c]["taxa"] = _taxa(por_cat[c]["entregas"])
    por_cat["SINISTRO"]["status"] = status_taxa(taxa_sin, "sinistro")
    por_cat["DEVOLUCAO"]["status"] = status_taxa(taxa_dev, "devolucao")
    por_cat["TENTATIVA_FRUSTRADA"]["status"] = status_taxa(taxa_tf, "tentativa_frustrada")

    return {
        "entregas": total,
        "afetadas": afetadas,
        "ocorrencias": ocorrencias_total,
        "valor_risco": round(valor_risco, 2),
        "taxa": taxa,
        "taxa_sinistro": taxa_sin,
        "taxa_devolucao": taxa_dev,
        "taxa_tentativa": taxa_tf,
        "status": status_taxa(taxa, "taxa_ocorrencia"),
        "por_categoria": por_cat,
    }


def _por_estado(rows: List[Dict]) -> List[Dict]:
    por_uf: Dict[str, List[Dict]] = {}
    for r in rows:
        if r["uf"]:
            por_uf.setdefault(r["uf"], []).append(r)
    linhas = []
    for uf in sorted(por_uf):
        ag = _agrega(por_uf[uf])
        ag["uf"] = uf
        linhas.append(ag)
    return linhas


def _serie_mensal(rows: List[Dict]) -> List[Dict]:
    por_mes: Dict[str, List[Dict]] = {}
    for r in rows:
        if r["mes"]:
            por_mes.setdefault(r["mes"], []).append(r)
    serie = []
    for m in sorted(por_mes):
        ag = _agrega(por_mes[m])
        serie.append({
            "mes": m, "taxa": ag["taxa"], "afetadas": ag["afetadas"],
            "entregas": ag["entregas"], "ocorrencias": ag["ocorrencias"],
            "valor_risco": ag["valor_risco"], "status": ag["status"],
        })
    return serie


def _top_motivos(rows: List[Dict], limit: int = 12) -> List[Dict]:
    """Códigos de ocorrência mais frequentes (só problemáticos), p/ o ranking de motivos."""
    agg: Dict[str, Dict] = {}
    for r in rows:
        for o in r["ocorrencias"]:
            c = o.get("categoria") or "OUTROS"
            if c not in PROBLEMA:
                continue
            cod = o.get("codigo") or "??"
            a = agg.setdefault(cod, {"codigo": cod, "descricao": o.get("descricao") or "",
                                     "categoria": c, "ocorrencias": 0})
            a["ocorrencias"] += 1
    linhas = sorted(agg.values(), key=lambda x: -x["ocorrencias"])
    return linhas[:limit]


# Categorias que caracterizam uma praça-problema (as que têm meta no painel):
# tentativa frustrada, sinistro (avaria/extravio) e devolução.
_CATS_PROBLEMA = set(_CAT_META.keys())


def reincidencia_por_cidade(mes: Optional[str] = None, uf: Optional[str] = None,
                            min_afetadas: int = 2) -> List[Dict]:
    """Cidades (destino) × transportadora com ocorrências problemáticas recorrentes.
    Para cada par, quantas entregas foram afetadas, quantas ocorrências, a taxa
    sobre as entregas da praça, o motivo dominante e o valor sob risco. Ranqueia as
    praças-problema (≥`min_afetadas` entregas afetadas) — base para renegociar prazo
    ou trocar de transportadora naquela região."""
    agg: Dict = {}
    for r in _universo(mes):
        if uf and r["uf"] != uf:
            continue
        destino = (r.get("destino") or "—").strip() or "—"
        key = (destino, r["transportadora"])
        a = agg.setdefault(key, {"destino": destino, "uf": r["uf"], "transportadora": r["transportadora"],
                                 "entregas_total": 0, "entregas_afetadas": 0, "ocorrencias": 0,
                                 "valor_risco": 0.0, "motivos": {}})
        a["entregas_total"] += 1
        probs = [o for o in r["ocorrencias"] if (o.get("categoria") in _CATS_PROBLEMA)]
        if not probs:
            continue
        a["entregas_afetadas"] += 1
        a["ocorrencias"] += len(probs)
        a["valor_risco"] += r.get("valor") or 0.0
        for o in probs:
            c = o.get("categoria") or "OUTROS"
            a["motivos"][c] = a["motivos"].get(c, 0) + 1
    out = []
    for a in agg.values():
        if a["entregas_afetadas"] < min_afetadas:
            continue
        motivo_top = max(a["motivos"].items(), key=lambda kv: kv[1])[0] if a["motivos"] else None
        out.append({
            "destino": a["destino"], "uf": a["uf"], "transportadora": a["transportadora"],
            "entregas_total": a["entregas_total"], "entregas_afetadas": a["entregas_afetadas"],
            "ocorrencias": a["ocorrencias"],
            "taxa_pct": round(100.0 * a["entregas_afetadas"] / a["entregas_total"], 1) if a["entregas_total"] else None,
            "valor_risco": round(a["valor_risco"], 2),
            "motivo_top": CAT_LABEL.get(motivo_top, motivo_top),
        })
    out.sort(key=lambda x: (-x["entregas_afetadas"], -x["ocorrencias"]))
    return out


def mapa(mes: Optional[str] = None, uf: Optional[str] = None) -> Dict:
    """Consolidado do Mapa de Ocorrências: resumo geral + por transportadora
    (com série mensal em cada) + série mensal geral + por estado + top motivos.
    `mes`='YYYY-MM'; `uf` filtra o painel (o mapa por estado mantém todas as UFs)."""
    universo = _universo(mes)

    # Mapa por estado é sempre sobre o mês inteiro (para permitir trocar de UF).
    por_estado = _por_estado(universo)

    uf = (uf or "").strip().upper() or None
    rows = [r for r in universo if r["uf"] == uf] if uf else universo

    geral = _agrega(rows)

    por_transp: Dict[str, List[Dict]] = {}
    for r in rows:
        por_transp.setdefault(r["transportadora"], []).append(r)

    linhas = []
    for nome, rs in por_transp.items():
        ag = _agrega(rs)
        ag["transportadora"] = nome
        ag["serie"] = _serie_mensal(rs)
        linhas.append(ag)

    # Pior primeiro: mais afetadas / maior taxa no topo (é o que exige ação).
    linhas.sort(key=lambda x: (-(x["taxa"] or 0), -x["afetadas"], -x["valor_risco"]))

    avaliadas = [l for l in linhas if l["taxa"] is not None]
    resumo = dict(geral)
    resumo.update({
        "transportadoras": len(linhas),
        "no_alvo": sum(1 for l in avaliadas if l["status"] == "verde"),
        "em_atencao": sum(1 for l in avaliadas if l["status"] == "amarelo"),
        "criticas": sum(1 for l in avaliadas if l["status"] == "vermelho"),
        "sem_vinculo": tracking.total_sem_vinculo(),
    })

    return {
        "metas": load_metas(),
        "mes": mes,
        "uf": uf,
        "resumo": resumo,
        "por_transportadora": linhas,
        "serie_mensal": _serie_mensal(rows),
        "por_estado": por_estado,
        "top_motivos": _top_motivos(rows),
        "cat_label": CAT_LABEL,
    }


def list_meses() -> List[str]:
    return audit_history.list_meses()


# ─── Exportação Excel ────────────────────────────────────────────────────────

def export_xlsx(data: Dict) -> bytes:
    """Relatório: Resumo + Por transportadora + Por estado + Top motivos + Evolução mensal."""
    import io
    from datetime import datetime
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from app.history import _sanitize_cell

    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    money_fmt = '"R$" #,##0.00'
    pct_fmt = "0.00"
    _situacao = {"verde": "No alvo", "amarelo": "Atenção", "vermelho": "Crítico", None: "—"}

    def _cabecalho(ws, headers, row):
        for col_idx, label in enumerate(headers, start=1):
            c = ws.cell(row=row, column=col_idx, value=label)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center")

    def _fill(status):
        cor = {"verde": "C6EFCE", "amarelo": "FFEB9C", "vermelho": "FFC7CE"}.get(status)
        return PatternFill(start_color=cor, end_color=cor, fill_type="solid") if cor else None

    r = data["resumo"]
    mes = data.get("mes")
    wb = Workbook()

    # ── Aba 1: Resumo ───────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = "Mapa de Ocorrências — não-conformidade nas entregas"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}" + (f" — Mês {mes}" if mes else " — Todos os meses")
    linhas = [
        ("Taxa de ocorrência (entregas afetadas)", r["taxa"], pct_fmt),
        ("Taxa de sinistro", r["taxa_sinistro"], pct_fmt),
        ("Taxa de devolução", r["taxa_devolucao"], pct_fmt),
        ("Taxa de tentativa frustrada", r["taxa_tentativa"], pct_fmt),
        ("R$ sob risco (carga das entregas afetadas)", r["valor_risco"], money_fmt),
        ("Entregas no período", r["entregas"], None),
        ("Entregas afetadas", r["afetadas"], None),
        ("Ocorrências problemáticas (eventos)", r["ocorrencias"], None),
        ("Transportadoras", r["transportadoras"], None),
        ("No alvo / Atenção / Críticas", f'{r["no_alvo"]} / {r["em_atencao"]} / {r["criticas"]}', None),
        ("Ocorrências sem vínculo (fora do mapa)", r["sem_vinculo"], None),
    ]
    for i, (label, valor, fmt) in enumerate(linhas, start=4):
        ws[f"A{i}"] = label
        ws[f"A{i}"].font = Font(bold=True)
        cell = ws[f"B{i}"]
        cell.value = valor if valor is not None else "—"
        if fmt and isinstance(valor, (int, float)):
            cell.number_format = fmt
    f = _fill(r["status"])
    if f:
        ws["B4"].fill = f
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 18

    # ── Aba 2: Por transportadora ───────────────────────────────────────
    ws2 = wb.create_sheet("Por transportadora")
    ws2["A1"] = "Ocorrências por transportadora"
    ws2["A1"].font = Font(bold=True, size=13)
    headers = ["Transportadora", "Taxa ocorr.", "Situação", "Sinistro %", "Devolução %",
               "Tent. frustr. %", "Entregas", "Afetadas", "Ocorrências", "R$ sob risco"]
    _cabecalho(ws2, headers, 3)
    for row_idx, t in enumerate(data["por_transportadora"], start=4):
        vals = [t["transportadora"], t["taxa"], _situacao.get(t["status"], "—"),
                t["taxa_sinistro"], t["taxa_devolucao"], t["taxa_tentativa"],
                t["entregas"], t["afetadas"], t["ocorrencias"], t["valor_risco"]]
        for col_idx, val in enumerate(vals, start=1):
            c = ws2.cell(row=row_idx, column=col_idx, value=_sanitize_cell(val if val is not None else "—"))
            if col_idx in (2, 4, 5, 6) and isinstance(val, (int, float)):
                c.number_format = pct_fmt
            elif col_idx == 10 and isinstance(val, (int, float)):
                c.number_format = money_fmt
        f = _fill(t["status"])
        if f:
            ws2.cell(row=row_idx, column=2).fill = f
            ws2.cell(row=row_idx, column=3).fill = f
    for col_idx, w in enumerate([26, 11, 10, 11, 12, 14, 9, 9, 11, 16], start=1):
        ws2.column_dimensions[get_column_letter(col_idx)].width = w
    ws2.freeze_panes = "A4"

    # ── Aba 3: Por estado ───────────────────────────────────────────────
    ws3 = wb.create_sheet("Por estado")
    ws3["A1"] = "Ocorrências por estado (UF) de destino"
    ws3["A1"].font = Font(bold=True, size=13)
    _cabecalho(ws3, ["UF", "Taxa ocorr.", "Situação", "Sinistro %", "Devolução %",
                     "Entregas", "Afetadas", "R$ sob risco"], 3)
    for row_idx, s in enumerate(data.get("por_estado", []), start=4):
        vals = [s["uf"], s["taxa"], _situacao.get(s["status"], "—"), s["taxa_sinistro"],
                s["taxa_devolucao"], s["entregas"], s["afetadas"], s["valor_risco"]]
        for col_idx, val in enumerate(vals, start=1):
            c = ws3.cell(row=row_idx, column=col_idx, value=_sanitize_cell(val if val is not None else "—"))
            if col_idx in (2, 4, 5) and isinstance(val, (int, float)):
                c.number_format = pct_fmt
            elif col_idx == 8 and isinstance(val, (int, float)):
                c.number_format = money_fmt
        f = _fill(s["status"])
        if f:
            ws3.cell(row=row_idx, column=2).fill = f
            ws3.cell(row=row_idx, column=3).fill = f
    for col_idx, w in enumerate([6, 11, 10, 11, 12, 9, 9, 16], start=1):
        ws3.column_dimensions[get_column_letter(col_idx)].width = w
    ws3.freeze_panes = "A4"

    # ── Aba 4: Top motivos ──────────────────────────────────────────────
    ws4 = wb.create_sheet("Top motivos")
    ws4["A1"] = "Motivos mais frequentes (código de ocorrência)"
    ws4["A1"].font = Font(bold=True, size=13)
    _cabecalho(ws4, ["Código", "Descrição", "Categoria", "Ocorrências"], 3)
    for row_idx, m in enumerate(data.get("top_motivos", []), start=4):
        vals = [m["codigo"], m["descricao"], CAT_LABEL.get(m["categoria"], m["categoria"]), m["ocorrencias"]]
        for col_idx, val in enumerate(vals, start=1):
            ws4.cell(row=row_idx, column=col_idx, value=_sanitize_cell(val))
    for col_idx, w in enumerate([9, 52, 26, 12], start=1):
        ws4.column_dimensions[get_column_letter(col_idx)].width = w
    ws4.freeze_panes = "A4"

    # ── Aba 5: Evolução mensal ──────────────────────────────────────────
    ws5 = wb.create_sheet("Evolução mensal")
    ws5["A1"] = "Evolução mensal da taxa de ocorrência"
    ws5["A1"].font = Font(bold=True, size=13)
    _cabecalho(ws5, ["Mês", "Taxa ocorr.", "Entregas", "Afetadas", "Ocorrências", "R$ sob risco"], 3)
    for row_idx, s in enumerate(data.get("serie_mensal", []), start=4):
        vals = [s["mes"], s["taxa"], s["entregas"], s["afetadas"], s["ocorrencias"], s["valor_risco"]]
        for col_idx, val in enumerate(vals, start=1):
            c = ws5.cell(row=row_idx, column=col_idx, value=_sanitize_cell(val if val is not None else "—"))
            if col_idx == 2 and isinstance(val, (int, float)):
                c.number_format = pct_fmt
            elif col_idx == 6 and isinstance(val, (int, float)):
                c.number_format = money_fmt
        f = _fill(s["status"])
        if f:
            ws5.cell(row=row_idx, column=2).fill = f
    for col_idx, w in enumerate([12, 11, 9, 9, 11, 16], start=1):
        ws5.column_dimensions[get_column_letter(col_idx)].width = w
    ws5.freeze_panes = "A4"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
