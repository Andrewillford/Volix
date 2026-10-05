"""
Coleta em Fornecedores — controle das coletas de mercadoria nos fornecedores.

MODELO (decidido com o usuário): uma COLETA é uma coletânea — várias NF-e coletadas
juntas (numa ou mais carretas), com UM frete total fechado. O usuário não sabe qual
NF foi em qual carreta, então não se mapeia NF↔carreta: o que importa é a coleta como
um todo (Σ das NFs) e o frete total.

  • Cabeçalho: data, fornecedor, transportadora, origem, destino, previsão, status,
    nº de carretas e valor do frete (total).
  • Notas: lista de NF-e (número, valor, produtos) — vinda da importação dos XMLs.
  • Indicador-chave: %frete = frete total ÷ Σ valor das NFs × 100.

Importação em lote: o operador sobe VÁRIOS XMLs de uma vez (`dados_de_xmls`), o painel
"explode" a lista das NFs somando os valores, e ele informa frete + previsão. Banco
próprio em SQLite (`tables/coletas.db`, env COLETAS_DB).
"""
from __future__ import annotations

import io
import json
import os
import re
import sqlite3
import unicodedata
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from app import nfe_parser

try:
    from app import cadastro
except Exception:  # pragma: no cover - defensivo; cadastro sempre existe
    cadastro = None

DB_PATH = os.getenv("COLETAS_DB", str(Path(__file__).parent.parent / "tables" / "coletas.db"))

STATUS_COLETA = ("Agendada", "Coletada", "Em trânsito", "Entregue", "Cancelada")
STATUS_DEFAULT = "Agendada"

# Campos escalares do cabeçalho da coleta (o resto — notas/valor_nf — é derivado).
_CAMPOS = (
    "data_coleta", "fornecedor", "valor_nf", "valor_frete",
    "transportadora", "origem", "destino", "previsao_entrega", "status", "observacoes",
)


# ─── Banco ───────────────────────────────────────────────────────────────────

def _get_conn() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db() -> None:
    Path(DB_PATH).parent.mkdir(parents=True, exist_ok=True)
    conn = _get_conn()
    conn.execute("""
        CREATE TABLE IF NOT EXISTS coletas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            data_coleta TEXT,
            fornecedor TEXT DEFAULT '',
            valor_nf REAL DEFAULT 0,
            valor_frete REAL DEFAULT 0,
            qtd_carretas INTEGER DEFAULT 1,
            transportadora TEXT DEFAULT '',
            origem TEXT DEFAULT '',
            destino TEXT DEFAULT '',
            previsao_entrega TEXT DEFAULT '',
            status TEXT DEFAULT 'Agendada',
            observacoes TEXT DEFAULT '',
            notas TEXT DEFAULT '',
            criado_em TEXT,
            atualizado_em TEXT
        )
    """)
    conn.execute("CREATE INDEX IF NOT EXISTS idx_coletas_data ON coletas(data_coleta)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_coletas_transp ON coletas(transportadora)")
    _migrar_colunas(conn)
    conn.commit()
    conn.close()


def _migrar_colunas(conn: sqlite3.Connection) -> None:
    """Adiciona colunas novas em bancos já criados (idempotente)."""
    cols = {r[1] for r in conn.execute("PRAGMA table_info(coletas)").fetchall()}
    for c, tipo in (("qtd_carretas", "INTEGER DEFAULT 1"), ("notas", "TEXT DEFAULT ''")):
        if c not in cols:
            conn.execute(f"ALTER TABLE coletas ADD COLUMN {c} {tipo}")


def limpar_tudo() -> int:
    """Apaga TODAS as coletas (não mexe em cadastro/usuários). Retorna quantas removeu."""
    conn = _get_conn()
    n = conn.execute("SELECT COUNT(*) FROM coletas").fetchone()[0]
    conn.execute("DELETE FROM coletas")
    conn.commit()
    conn.close()
    return n


def _now() -> str:
    return datetime.now().isoformat(timespec="seconds")


# ─── Helpers ─────────────────────────────────────────────────────────────────

def _num(v) -> float:
    """Converte para float tolerando string com vírgula/moeda; falha vira 0."""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s:
        return 0.0
    s = s.replace("R$", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _int(v, default: int = 0) -> int:
    try:
        return int(round(_num(v))) if v not in (None, "") else default
    except (ValueError, TypeError):
        return default


def _norm(s) -> str:
    """Normaliza texto p/ casar status: sem acento, minúsculo, só a-z0-9."""
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]", "", s.lower())


def _pct_frete(valor_frete: float, valor_nf: float) -> Optional[float]:
    """%frete = frete total ÷ Σ valor das NFs × 100. Sem NF → None (fora da média)."""
    if valor_nf and valor_nf > 0:
        return round(valor_frete / valor_nf * 100, 2)
    return None


def _limpa_status(status: Optional[str]) -> str:
    s = _norm(status)
    if not s:
        return STATUS_DEFAULT
    for opt in STATUS_COLETA:
        if s == _norm(opt):
            return opt
    return STATUS_DEFAULT


def _parse_json_list(raw) -> List[Dict]:
    if not raw:
        return []
    try:
        v = json.loads(raw)
        return v if isinstance(v, list) else []
    except (ValueError, TypeError):
        return []


def _norm_nota(n: Dict) -> Dict:
    """Padroniza uma NF da coleta (o que guardamos por nota)."""
    return {
        "chave": re.sub(r"\D", "", str(n.get("chave") or "")),
        "numero": str(n.get("numero") if n.get("numero") is not None else n.get("nf_numero") or "").strip(),
        "fornecedor": str(n.get("fornecedor") or n.get("emitente") or "").strip(),
        "valor": round(_num(n.get("valor") if n.get("valor") is not None else n.get("valor_nf")), 2),
        "itens": n.get("itens") if isinstance(n.get("itens"), list) else [],
    }


def _prep_notas(data: Dict) -> Optional[List[Dict]]:
    """Extrai/normaliza as notas do payload. None = não foi enviado (não mexer)."""
    raw = data.get("notas")
    if not isinstance(raw, list):
        return None
    return [_norm_nota(n) for n in raw]


def _fornecedor_comum(notas: List[Dict]) -> str:
    nomes = [n.get("fornecedor", "").strip() for n in notas if n.get("fornecedor")]
    return Counter(nomes).most_common(1)[0][0] if nomes else ""


def _row_para_dict(row: sqlite3.Row) -> Dict:
    d = {k: row[k] for k in row.keys()}
    notas = _parse_json_list(d.get("notas"))
    d["notas"] = notas
    d["qtd_nfs"] = len(notas)
    d["qtd_itens"] = sum(len(n.get("itens") or []) for n in notas)
    d["valor_nf"] = _num(d.get("valor_nf"))
    d["valor_frete"] = _num(d.get("valor_frete"))
    d["qtd_carretas"] = int(d.get("qtd_carretas") or 1)
    d["pct_frete"] = _pct_frete(d["valor_frete"], d["valor_nf"])
    return d


def _normaliza_entrada(data: Dict) -> Dict:
    """Mapeia o payload para as colunas escalares, limpando números/status/datas."""
    out = {c: (data.get(c) if data.get(c) is not None else "") for c in _CAMPOS}
    out["valor_nf"] = _num(data.get("valor_nf"))
    out["valor_frete"] = _num(data.get("valor_frete"))
    out["status"] = _limpa_status(data.get("status"))
    for c in ("fornecedor", "transportadora", "origem", "destino", "observacoes"):
        out[c] = str(out[c]).strip()
    for c in ("data_coleta", "previsao_entrega"):
        v = str(out[c]).strip()
        out[c] = v[:10] if len(v) >= 10 else v
    return out


# ─── CRUD ────────────────────────────────────────────────────────────────────

def list_coletas(mes: Optional[str] = None, status: Optional[str] = None,
                 transportadora: Optional[str] = None, search: Optional[str] = None) -> List[Dict]:
    conn = _get_conn()
    where, params = [], []
    if mes:
        where.append("substr(data_coleta,1,7)=?"); params.append(mes)
    if status:
        where.append("status=?"); params.append(status)
    if transportadora:
        where.append("transportadora=?"); params.append(transportadora)
    if search:
        like = f"%{search.strip()}%"
        where.append("(fornecedor LIKE ? OR origem LIKE ? OR destino LIKE ? OR transportadora LIKE ? OR notas LIKE ?)")
        params.extend([like, like, like, like, like])
    sql = "SELECT * FROM coletas"
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY date(data_coleta) DESC, id DESC"
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    return [_row_para_dict(r) for r in rows]


def get_coleta(cid: int) -> Optional[Dict]:
    conn = _get_conn()
    row = conn.execute("SELECT * FROM coletas WHERE id=?", (cid,)).fetchone()
    conn.close()
    return _row_para_dict(row) if row else None


def _monta_valores(data: Dict, notas: Optional[List[Dict]], base: Optional[Dict] = None) -> Tuple[Dict, List[Dict]]:
    """Resolve os campos escalares + notas efetivas. `base` = coleta atual (no update)."""
    origem = dict(base) if base else {}
    merged = {c: origem.get(c) for c in _CAMPOS}
    for c in _CAMPOS:
        if c in data and data[c] is not None:
            merged[c] = data[c]
    notas_efet = notas if notas is not None else (origem.get("notas") or [])
    if notas_efet:
        merged["valor_nf"] = round(sum(_num(n.get("valor")) for n in notas_efet), 2)
        if not str(merged.get("fornecedor") or "").strip():
            merged["fornecedor"] = _fornecedor_comum(notas_efet)
    return _normaliza_entrada(merged), notas_efet


def create_coleta(data: Dict) -> Dict:
    notas = _prep_notas(data)
    d, notas_efet = _monta_valores(data, notas)
    qtd_carretas = max(1, _int(data.get("qtd_carretas"), 1))
    notas_json = json.dumps(notas_efet, ensure_ascii=False)
    conn = _get_conn()
    cols = list(_CAMPOS) + ["qtd_carretas", "notas", "criado_em", "atualizado_em"]
    vals = [d[c] for c in _CAMPOS] + [qtd_carretas, notas_json, _now(), _now()]
    cur = conn.execute(
        f"INSERT INTO coletas ({', '.join(cols)}) VALUES ({', '.join(['?'] * len(cols))})", vals,
    )
    conn.commit()
    cid = cur.lastrowid
    conn.close()
    return get_coleta(cid)


def update_coleta(cid: int, data: Dict) -> Optional[Dict]:
    atual = get_coleta(cid)
    if not atual:
        return None
    notas = _prep_notas(data)  # None = mantém as atuais
    d, notas_efet = _monta_valores(data, notas, base=atual)
    qtd_carretas = max(1, _int(data.get("qtd_carretas"), atual.get("qtd_carretas") or 1))
    notas_json = json.dumps(notas_efet, ensure_ascii=False)
    conn = _get_conn()
    campos = list(_CAMPOS) + ["qtd_carretas", "notas"]
    valores = [d[c] for c in _CAMPOS] + [qtd_carretas, notas_json]
    sets = ", ".join(f"{c}=?" for c in campos) + ", atualizado_em=?"
    conn.execute(f"UPDATE coletas SET {sets} WHERE id=?", valores + [_now(), cid])
    conn.commit()
    conn.close()
    return get_coleta(cid)


def delete_coleta(cid: int) -> bool:
    conn = _get_conn()
    cur = conn.execute("DELETE FROM coletas WHERE id=?", (cid,))
    conn.commit()
    n = cur.rowcount
    conn.close()
    return n > 0


# ─── Importação de XMLs (lote) ───────────────────────────────────────────────

def _cidade_uf(municipio: str, uf: str) -> str:
    municipio = (municipio or "").strip()
    uf = (uf or "").strip().upper()
    if municipio and uf:
        return f"{municipio}/{uf}"
    return municipio or uf


def _resolve_transportadora(cnpj: str, nome_nf: str) -> str:
    cc = re.sub(r"\D", "", cnpj or "")
    if cc and cadastro is not None:
        try:
            for t in cadastro.list_transportadoras():
                for c in (t.get("cnpjs") or []):
                    if re.sub(r"\D", "", c or "") == cc:
                        return t.get("nome") or (nome_nf or "").strip()
        except Exception:
            pass
    return (nome_nf or "").strip()


def chave_existe(chave: Optional[str]) -> bool:
    """Alguma coleta já contém uma NF com esta chave (44 díg)? (busca no JSON notas)."""
    ch = re.sub(r"\D", "", str(chave or ""))
    if len(ch) != 44:
        return False
    conn = _get_conn()
    row = conn.execute("SELECT 1 FROM coletas WHERE notas LIKE ? LIMIT 1", (f"%{ch}%",)).fetchone()
    conn.close()
    return row is not None


def dados_de_xmls(files: List[Tuple[str, bytes]]) -> Dict:
    """Lê VÁRIOS XMLs de NF-e e devolve a lista de notas + um cabeçalho comum
    (fornecedor/origem/destino/transportadora/data) para preencher a coleta. Não grava."""
    notas: List[Dict] = []
    erros: List[Dict] = []
    for filename, content in files:
        try:
            nf = nfe_parser.parse_nfe_bytes(content, filename)
        except nfe_parser.NfeParseError as e:
            erros.append({"arquivo": filename, "erro": str(e)})
            continue
        notas.append({
            "arquivo": filename,
            "chave": nf.chave,
            "numero": nf.numero,
            "fornecedor": nf.emitente_nome,
            "valor": round(nf.valor_nf or 0, 2),
            "origem": _cidade_uf(nf.emitente_municipio, nf.emitente_uf),
            "destino": _cidade_uf(nf.dest_municipio, nf.dest_uf),
            "transportadora": _resolve_transportadora(nf.transportadora_cnpj, nf.transportadora_nome),
            "data_emissao": (nf.data_emissao or "")[:10] if nf.data_emissao else "",
            "itens": nf.itens,
            "qtd_itens": len(nf.itens),
            "ja_existe": chave_existe(nf.chave),
        })

    def comum(campo: str) -> str:
        vals = [n[campo] for n in notas if n.get(campo)]
        return Counter(vals).most_common(1)[0][0] if vals else ""

    datas = [n["data_emissao"] for n in notas if n.get("data_emissao")]
    return {
        "notas": notas,
        "erros": erros,
        "qtd_nfs": len(notas),
        "qtd_itens": sum(n["qtd_itens"] for n in notas),
        "valor_nf_total": round(sum(n["valor"] for n in notas), 2),
        "fornecedor": comum("fornecedor"),
        "origem": comum("origem"),
        "destino": comum("destino"),
        "transportadora": comum("transportadora"),
        "data_coleta": min(datas) if datas else datetime.now().strftime("%Y-%m-%d"),
        "duplicadas": [n["numero"] for n in notas if n.get("ja_existe")],
    }


# ─── Agregações (KPIs + gráficos) ────────────────────────────────────────────

def _agrega(coletas: List[Dict]) -> Dict:
    total = len(coletas)
    frete_total = sum(_num(c.get("valor_frete")) for c in coletas)
    com_nf = [c for c in coletas if _num(c.get("valor_nf")) > 0]
    valor_nf = sum(_num(c.get("valor_nf")) for c in com_nf)
    frete_base = sum(_num(c.get("valor_frete")) for c in com_nf)
    pct = round(frete_base / valor_nf * 100, 2) if valor_nf > 0 else None
    pcts = [c["pct_frete"] for c in com_nf if c.get("pct_frete") is not None]
    pct_simples = round(sum(pcts) / len(pcts), 2) if pcts else None
    return {
        "coletas": total,
        "com_valor_nf": len(com_nf),
        "sem_valor_nf": total - len(com_nf),
        "nfs": sum(c.get("qtd_nfs", 0) for c in coletas),
        "carretas": sum(c.get("qtd_carretas", 1) for c in coletas),
        "frete_total": round(frete_total, 2),
        "valor_nf": round(valor_nf, 2),
        "custo_medio_pct": pct,
        "custo_medio_pct_simples": pct_simples,
    }


def _serie_mensal(coletas: List[Dict]) -> List[Dict]:
    buckets: Dict[str, Dict[str, float]] = {}
    for c in coletas:
        mes = (c.get("data_coleta") or "")[:7]
        if len(mes) != 7:
            continue
        nf = _num(c.get("valor_nf"))
        if nf <= 0:
            continue
        b = buckets.setdefault(mes, {"frete": 0.0, "valor_nf": 0.0, "n": 0})
        b["frete"] += _num(c.get("valor_frete"))
        b["valor_nf"] += nf
        b["n"] += 1
    saida = []
    for mes in sorted(buckets):
        b = buckets[mes]
        pct = round(b["frete"] / b["valor_nf"] * 100, 2) if b["valor_nf"] > 0 else None
        saida.append({"mes": mes, "frete": round(b["frete"], 2),
                      "valor_nf": round(b["valor_nf"], 2), "coletas": b["n"], "pct_frete": pct})
    return saida


def _por_transportadora(coletas: List[Dict]) -> List[Dict]:
    buckets: Dict[str, Dict[str, float]] = {}
    for c in coletas:
        nome = (c.get("transportadora") or "").strip() or "— sem transportadora —"
        b = buckets.setdefault(nome, {"frete": 0.0, "valor_nf": 0.0, "n": 0})
        b["frete"] += _num(c.get("valor_frete"))
        b["valor_nf"] += _num(c.get("valor_nf"))
        b["n"] += 1
    saida = []
    for nome, b in buckets.items():
        pct = round(b["frete"] / b["valor_nf"] * 100, 2) if b["valor_nf"] > 0 else None
        saida.append({"transportadora": nome, "frete": round(b["frete"], 2),
                      "valor_nf": round(b["valor_nf"], 2), "coletas": b["n"], "pct_frete": pct})
    saida.sort(key=lambda x: (x["pct_frete"] is None, -(x["pct_frete"] or 0)))
    return saida


def meses_disponiveis() -> List[str]:
    conn = _get_conn()
    rows = conn.execute(
        "SELECT DISTINCT substr(data_coleta,1,7) AS mes FROM coletas "
        "WHERE length(data_coleta)>=7 ORDER BY mes DESC"
    ).fetchall()
    conn.close()
    return [r["mes"] for r in rows if r["mes"]]


def _nomes_transportadoras() -> List[str]:
    if cadastro is None:
        return []
    try:
        return sorted({(t.get("nome") or "").strip() for t in cadastro.list_transportadoras() if t.get("nome")})
    except Exception:
        return []


def painel(mes: Optional[str] = None, status: Optional[str] = None,
           transportadora: Optional[str] = None, search: Optional[str] = None,
           limit: Optional[int] = None, offset: int = 0) -> Dict:
    """Painel de coletas. Os KPIs, séries e o Excel são sempre sobre TODO o
    conjunto filtrado; apenas a lista da tabela (`coletas`) é paginada quando
    `limit` é informado, para não despejar centenas de linhas no DOM. `total` é
    a contagem do conjunto filtrado inteiro (base do pager)."""
    coletas = list_coletas(mes=mes, status=status, transportadora=transportadora, search=search)
    total = len(coletas)
    if limit is not None:
        offset = max(0, int(offset or 0))
        pagina = coletas[offset:offset + max(1, int(limit))]
    else:
        pagina = coletas
    return {
        "coletas": pagina,
        "total": total,
        "resumo": _agrega(coletas),
        "serie_mensal": _serie_mensal(coletas),
        "por_transportadora": _por_transportadora(coletas),
        "meses": meses_disponiveis(),
        "status_opcoes": list(STATUS_COLETA),
        "transportadoras": _nomes_transportadoras(),
    }


# ─── Exportação (relatório Excel) ────────────────────────────────────────────

def export_xlsx(mes: Optional[str] = None, status: Optional[str] = None,
                transportadora: Optional[str] = None, search: Optional[str] = None) -> bytes:
    """Relatório Excel: aba Coletas (uma linha por coleta) + aba NFs (uma por nota) +
    aba Resumo (KPIs, por transportadora, por mês). Respeita os filtros do painel."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    dados = painel(mes=mes, status=status, transportadora=transportadora, search=search)
    coletas = dados["coletas"]
    resumo = dados["resumo"]

    RS = '"R$" #,##0.00'
    PCT = '0.00"%"'
    head_fill = PatternFill("solid", fgColor="1F2A44")
    head_font = Font(bold=True, color="FFFFFF")

    def cabecalho(ws, ncols, row=1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 24

    wb = openpyxl.Workbook()

    # ── Aba Coletas ──
    ws = wb.active
    ws.title = "Coletas"
    cols = ["ID", "Data", "Fornecedor", "NFs", "Valor NFs (R$)", "Carretas", "Frete total (R$)",
            "% Frete", "Transportadora", "Origem", "Destino", "Previsão", "Status", "Observações"]
    ws.append(cols)
    cabecalho(ws, len(cols))
    for c in coletas:
        ws.append([
            c.get("id"), c.get("data_coleta") or "", c.get("fornecedor") or "",
            c.get("qtd_nfs", 0), _num(c.get("valor_nf")), c.get("qtd_carretas", 1),
            _num(c.get("valor_frete")), c.get("pct_frete"),
            c.get("transportadora") or "", c.get("origem") or "", c.get("destino") or "",
            c.get("previsao_entrega") or "", c.get("status") or "", c.get("observacoes") or "",
        ])
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=5).number_format = RS
        ws.cell(row=r, column=7).number_format = RS
        ws.cell(row=r, column=8).number_format = PCT
    for col, w in zip("ABCDEFGHIJKLMN", (5, 12, 28, 6, 15, 8, 15, 9, 18, 18, 18, 12, 12, 26), strict=False):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    if not coletas:
        ws.cell(row=2, column=1, value="Nenhuma coleta lançada para os filtros atuais.")

    # ── Aba NFs (uma linha por nota, com a coleta a que pertence) ──
    wn = wb.create_sheet("NFs")
    ncols = ["Coleta ID", "Data", "Fornecedor (coleta)", "Nº NF", "Fornecedor (NF)",
             "Valor NF (R$)", "Produtos", "Transportadora"]
    wn.append(ncols)
    cabecalho(wn, len(ncols))
    for c in coletas:
        for n in c.get("notas", []):
            wn.append([
                c.get("id"), c.get("data_coleta") or "", c.get("fornecedor") or "",
                n.get("numero") or "", n.get("fornecedor") or "",
                _num(n.get("valor")), len(n.get("itens") or []), c.get("transportadora") or "",
            ])
    for r in range(2, wn.max_row + 1):
        wn.cell(row=r, column=6).number_format = RS
    for col, w in zip("ABCDEFGH", (9, 12, 26, 12, 26, 15, 9, 18), strict=False):
        wn.column_dimensions[col].width = w
    wn.freeze_panes = "A2"

    # ── Aba Resumo ──
    wr = wb.create_sheet("Resumo")
    wr.cell(row=1, column=1, value="RESUMO — COLETA EM FORNECEDORES").font = Font(bold=True, size=13, color="2E7D32")
    kpis = [
        ("Custo médio de frete (%NF)", resumo.get("custo_medio_pct"), PCT),
        ("Custo médio (média simples)", resumo.get("custo_medio_pct_simples"), PCT),
        ("Coletas", resumo.get("coletas"), "0"),
        ("Notas fiscais", resumo.get("nfs"), "0"),
        ("Carretas", resumo.get("carretas"), "0"),
        ("Frete total", resumo.get("frete_total"), RS),
        ("Valor das NFs", resumo.get("valor_nf"), RS),
    ]
    r = 3
    for label, val, fmt in kpis:
        wr.cell(row=r, column=1, value=label).font = Font(bold=True)
        wr.cell(row=r, column=2, value=val).number_format = fmt
        r += 1

    def bloco(titulo, cab, linhas, chave0):
        nonlocal r
        r += 1
        wr.cell(row=r, column=1, value=titulo).font = Font(bold=True)
        r += 1
        for i, t in enumerate(cab, start=1):
            wr.cell(row=r, column=i, value=t)
        cabecalho(wr, len(cab), row=r)
        r += 1
        for ln in linhas:
            wr.cell(row=r, column=1, value=ln.get(chave0))
            wr.cell(row=r, column=2, value=ln.get("coletas"))
            wr.cell(row=r, column=3, value=ln.get("frete")).number_format = RS
            wr.cell(row=r, column=4, value=ln.get("valor_nf")).number_format = RS
            wr.cell(row=r, column=5, value=ln.get("pct_frete")).number_format = PCT
            r += 1

    bloco("POR TRANSPORTADORA", ["Transportadora", "Coletas", "Frete (R$)", "Valor NF (R$)", "% Frete"],
          dados["por_transportadora"], "transportadora")
    bloco("POR MÊS", ["Mês", "Coletas", "Frete (R$)", "Valor NF (R$)", "% Frete"],
          dados["serie_mensal"], "mes")
    for col, w in zip("ABCDE", (30, 12, 15, 15, 10), strict=False):
        wr.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
